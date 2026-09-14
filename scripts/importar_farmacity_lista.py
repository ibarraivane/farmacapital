#!/usr/bin/env python3
"""Cruza la lista Farma City (Excel) → producto_precios_referencia.

Solo matches por código de barras contra productos que ya existen.
Usa PRECIO NETO (mayoreo). No inventa match por nombre.

  python3 scripts/importar_farmacity_lista.py --archivo lista.xlsx --sql-only
"""
from __future__ import annotations

import argparse
import os
import re
import sys
from pathlib import Path

import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "sql"
FECHA_LISTA = "2026-09-04"
FUENTE = "farmacity"
HOJA = "Hoja1"
ARCHIVO_NOMBRE = "MEDICAMENTO 4 DE SEPT2026.xlsx"


def cargar_env() -> dict[str, str]:
    out: dict[str, str] = {}
    for env_path in (Path("/tmp/fc-sb.env"), ROOT / ".env"):
        if not env_path.exists():
            continue
        for line in env_path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, _, v = line.partition("=")
            out.setdefault(k.strip(), v.strip().strip('"').strip("'"))
    out.update({k: v for k, v in os.environ.items() if v})
    return out


def digits(val) -> str:
    return re.sub(r"\D", "", str(val or ""))


def barcode_keys(code: str) -> set[str]:
    d = digits(code)
    if not d:
        return set()
    keys = {d, d.lstrip("0") or "0"}
    if len(d) < 12:
        keys.add(d.zfill(12))
    if len(d) < 13:
        keys.add(d.zfill(13))
    if len(d) == 12:
        keys.add("0" + d)
    if len(d) == 13 and d.startswith("0"):
        keys.add(d[1:])
    return {k for k in keys if k}


def fetch_productos(url: str, key: str) -> list[dict]:
    headers = {"apikey": key, "Authorization": f"Bearer {key}"}
    rows: list[dict] = []
    offset = 0
    while True:
        h = {**headers, "Range": f"{offset}-{offset + 499}"}
        r = requests.get(
            f"{url}/rest/v1/productos",
            headers=h,
            params={
                "select": "id,sku,nombre,codigo_barras,activo,costo",
                "activo": "eq.true",
                "order": "id.asc",
            },
            timeout=60,
        )
        r.raise_for_status()
        batch = r.json()
        rows.extend(batch)
        if len(batch) < 500:
            break
        offset += 500
    return rows


def index_por_ean(productos: list[dict]) -> dict[str, list[dict]]:
    idx: dict[str, list[dict]] = {}
    for p in productos:
        for key in barcode_keys(p.get("codigo_barras") or ""):
            idx.setdefault(key, []).append(p)
    return idx


def leer_lista(path: Path) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if HOJA not in wb.sheetnames:
        sys.exit(f"No está la hoja {HOJA!r}. Hojas: {wb.sheetnames}")
    ws = wb[HOJA]
    out: list[dict] = []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1 or not row:
            continue
        codigo, nombre = row[0], row[1]
        precio_neto = row[2] if len(row) > 2 else None
        exist = row[3] if len(row) > 3 else None
        ean = digits(codigo)
        if not ean:
            continue
        try:
            precio = float(precio_neto)
        except (TypeError, ValueError):
            continue
        if precio <= 0:
            continue
        nom = str(nombre or "").replace("_x000D_", " ").replace("\n", " ").strip()
        nom = re.sub(r"\s+", " ", nom)
        try:
            exist_n = int(float(exist)) if exist not in (None, "") else None
        except (TypeError, ValueError):
            exist_n = None
        out.append({
            "ean": ean,
            "nombre": nom,
            "precio": round(precio, 2),
            "exist": exist_n,
            "line": i,
        })
    wb.close()
    return out


def matchear(lista: list[dict], idx: dict[str, list[dict]]) -> tuple[list[dict], list[dict], int]:
    matched: list[dict] = []
    unmatched: list[dict] = []
    ambiguos = 0
    seen_producto: set[int] = set()
    for row in lista:
        hits: list[dict] = []
        seen: set[int] = set()
        for key in barcode_keys(row["ean"]):
            for p in idx.get(key, []):
                if p["id"] not in seen:
                    seen.add(p["id"])
                    hits.append(p)
        if len(hits) != 1:
            if len(hits) > 1:
                ambiguos += 1
            unmatched.append(row)
            continue
        prod = hits[0]
        if prod["id"] in seen_producto:
            continue
        seen_producto.add(prod["id"])
        matched.append({
            "producto_id": prod["id"],
            "sku": prod.get("sku"),
            "nombre_catalogo": prod.get("nombre"),
            "costo": prod.get("costo"),
            "ean": row["ean"],
            "nombre_fuente": row["nombre"],
            "precio": row["precio"],
            "exist": row.get("exist"),
        })
    return matched, unmatched, ambiguos


def sql_quote(s: str) -> str:
    return "'" + str(s).replace("'", "''") + "'"


def generate_sql(matched: list[dict], archivo: str, stats: dict) -> str:
    lines = [
        f"-- Farma City lista {FECHA_LISTA} · {len(matched)} matches EAN · precio neto",
        f"-- Archivo: {archivo}",
        f"-- Lista: {stats['lista']} renglones · catálogo: {stats['catalogo']} activos"
        f" · sin match: {stats['unmatched']} · EAN ambiguo: {stats['ambiguos']}",
        "-- Solo EAN exacto (variantes de ceros). No hay match por nombre.",
        "",
        "begin;",
        "",
        "insert into public.fuentes_precio (id, nombre, tipo, metodo, notas) values",
        "  ('farmacity', 'Farma City', 'compra', 'import_archivo',",
        "   'Lista Cityfarma Iztapalapa. Precio neto de mayoreo.')",
        "on conflict (id) do update set",
        "  nombre = excluded.nombre,",
        "  tipo = excluded.tipo,",
        "  metodo = excluded.metodo,",
        "  notas = excluded.notas;",
        "",
        "with imp as (",
        "  insert into public.importaciones_referencia (fuente, tipo, fecha_lista, archivo, filas_ok, notas)",
        f"  values ('farmacity', 'compra', '{FECHA_LISTA}', {sql_quote(archivo)}, {len(matched)},",
        "          'importar_farmacity_lista.py · precio neto · match EAN')",
        "  returning id",
        ")",
        "insert into public.producto_precios_referencia (",
        "  producto_id, fuente, tipo, precio, fecha, nombre_fuente, sku_externo,",
        "  confianza, origen, import_id, notas",
        ")",
        "select",
        f"  v.producto_id, 'farmacity', 'compra', v.precio, '{FECHA_LISTA}'::date,",
        "  v.nombre_fuente, v.sku_externo, 100, 'import_csv', imp.id,",
        "  'lista 4-sep-2026 · precio neto'",
        "from imp, (values",
    ]
    value_rows = []
    for m in matched:
        value_rows.append(
            f"  ({m['producto_id']}::bigint, {m['precio']}::numeric, "
            f"{sql_quote(m['nombre_fuente'])}, {sql_quote(m['ean'])})"
        )
    lines.append(",\n".join(value_rows))
    lines.extend([
        ") as v(producto_id, precio, nombre_fuente, sku_externo);",
        "",
        "commit;",
        "",
    ])
    return "\n".join(lines)


def main() -> int:
    parser = argparse.ArgumentParser(description="Importar lista Farma City a referencias de compra")
    parser.add_argument("--archivo", type=Path, required=True)
    parser.add_argument("--sql-only", action="store_true")
    args = parser.parse_args()

    if not args.archivo.exists():
        sys.exit(f"No existe: {args.archivo}")

    env = cargar_env()
    url = env.get("REACT_APP_SUPABASE_URL") or env.get("SUPABASE_URL") or ""
    key = (
        env.get("SUPABASE_SERVICE_ROLE_KEY")
        or env.get("REACT_APP_SUPABASE_ANON_KEY")
        or ""
    )
    if not url or not key:
        sys.exit("Falta URL/key de Supabase en .env")

    print("Catálogo Supabase…")
    productos = fetch_productos(url, key)
    print(f"  {len(productos)} activos")
    idx = index_por_ean(productos)

    print(f"Leyendo {args.archivo.name}…")
    lista = leer_lista(args.archivo)
    print(f"  {len(lista)} renglones con precio neto y EAN")

    matched, unmatched, ambiguos = matchear(lista, idx)
    mas_barato = 0
    for m in matched:
        try:
            costo = float(m["costo"]) if m.get("costo") is not None else None
        except (TypeError, ValueError):
            costo = None
        if costo and m["precio"] < costo - 0.05:
            mas_barato += 1

    print(f"Matches EAN: {len(matched)}")
    print(f"Sin match (no están en catálogo o EAN ambiguo): {len(unmatched)}")
    print(f"EAN ambiguo (más de un SKU): {ambiguos}")
    print(f"Más baratos que tu costo: {mas_barato}")
    print("Ejemplos:")
    for m in matched[:12]:
        print(f"  {m['sku']}  ${m['precio']:.2f}  {m['nombre_catalogo'][:48]}")

    stats = {
        "lista": len(lista),
        "catalogo": len(productos),
        "unmatched": len(unmatched),
        "ambiguos": ambiguos,
    }
    out_sql = OUT_DIR / f"patch_farmacity_lista_{FECHA_LISTA.replace('-', '')}.sql"
    out_sql.write_text(generate_sql(matched, ARCHIVO_NOMBRE, stats), encoding="utf-8")
    print(f"SQL: {out_sql}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
