#!/usr/bin/env python3
"""Cruza el surtido de Farmacias Similares contra el inventario vivo de FarmaCapital.

El Excel original (`pricing/fuentes/articulos_farmacias.xlsx`) no se versiona.
Si está en disco se usa; si no, se baja el catálogo público VTEX
(https://www.farmaciasdesimilares.com) — la misma lista de artículos de farmacia.

Compara por genérico (principio + concentración + forma), no por marca comercial:
si Similares vende ibuprofeno 400 mg 10 tabletas y nosotros tenemos AMSA/Ultra
del mismo PA/dosis/forma, cuenta como cubierto.

Uso:
  python3 scripts/cruce_similares_inventario.py
  python3 scripts/cruce_similares_inventario.py --excel pricing/fuentes/articulos_farmacias.xlsx
"""
from __future__ import annotations

import argparse
import csv
import importlib.util
import json
import re
import sys
import time
import unicodedata
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
EXCEL_DEFAULT = ROOT / "pricing" / "fuentes" / "articulos_farmacias.xlsx"
VTEX_SEARCH = "https://www.farmaciasdesimilares.com/api/catalog_system/pub/products/search"
HOY = date.today().isoformat()
REPORTE_DIR = ROOT / "pricing" / "reportes"
SNAPSHOT_DIR = ROOT / "pricing" / "importados"

spec = importlib.util.spec_from_file_location(
    "sync_excel", ROOT / "scripts" / "sync_precios_excel_farmacias.py"
)
ex = importlib.util.module_from_spec(spec)
spec.loader.exec_module(ex)
sync = ex.sync
sync.SINONIMOS.setdefault("cla", "clavulanico")
sync.SINONIMOS.setdefault("clavulanico", "clavulanico")
sync.SINONIMOS.setdefault("tioctico", "tioctico")
sync.SINONIMOS.setdefault("butilhio", "butilhioscina")
sync.SINONIMOS.setdefault("iodopovidona", "yodopovidona")
sync.SINONIMOS.setdefault("yodopovidona", "yodopovidona")
sync.SINONIMOS.setdefault("povidona", "yodopovidona")
sync.SINONIMOS.setdefault("simeticona", "dimeticona")
sync.SINONIMOS.setdefault("dimeticona", "dimeticona")
sync.SINONIMOS.setdefault("peroxido", "oxigenada")
sync.SINONIMOS.setdefault("oxigenada", "oxigenada")
sync.SINONIMOS.setdefault("mebendazol", "albendazol")
sync.SINONIMOS.setdefault("albendazol", "albendazol")

UMBRAL_CUBIERTO = 80
UMBRAL_EQUIV = 70

CLASE_A = {
    "paracetamol", "ibuprofeno", "naproxeno", "diclofenaco", "ketorolaco", "metamizol",
    "omeprazol", "pantoprazol", "ranitidina",
    "metformina", "glibenclamida",
    "losartan", "amlodipino", "enalapril", "captopril",
    "amoxicilina", "ampicilina", "ciprofloxacino", "metronidazol", "trimetoprima",
    "azitromicina",
    "loratadina", "cetirizina", "clorfenamina",
    "atorvastatina",
    "acetilsalicilico",
    "salbutamol", "ambroxol",
    "clotrimazol",
    "furosemida",
    "sildenafil",
}

MARCA_PROPIA_SIMI = re.compile(
    r"\b(simi\s?flex|simi\s?fx|simiflex|simibacil|simipro|simifibra|simifila|"
    r"simialoe|dr simi|multigomi|gomitas|nopal/chia|pick up)\b",
    re.I,
)

SKIP_CAT_RE = re.compile(
    r"souvenir|perfumer|alimento|promocional|materiales diversos|juguete",
    re.I,
)
INCLUDE_CAT_RE = re.compile(
    r"medicamento|curaci[oó]n|suplemento|salud sexual|maternidad|herbolar|"
    r"vitamin|diabetes|cardio|analg[eé]sic|antibi[oó]tic|respirator|"
    r"antihist|estomac|gastro|dermato|oftal|otic|ginecol|urolog|"
    r"material de curaci|botiqu[ií]n|primeros auxilios",
    re.I,
)

FORMA_FAMILIA = {
    "tableta": "oral_solido",
    "capsula": "oral_solido",
    "comprimido": "oral_solido",
    "gragea": "oral_solido",
    "suspension": "oral_liq",
    "jarabe": "oral_liq",
    "solucion": "oral_liq",
    "polvo": "oral_liq",
    "crema": "topico",
    "gel": "topico",
    "unguento": "topico",
    "pomada": "topico",
}

MARCAS_CASA_SIMI = {
    "pick up", "generico", "genérico", "farmacias similares", "similares",
    "dr simi", "simi",
}

# Palabras que aparecen en el nombre pero no identifican al fármaco.
# Sin esto, "AGUA DE MAR" pegaba con "Agua oxigenada" por el token "agua".
PA_DEBIL = {
    "agua", "mar", "gel", "crema", "solucion", "tableta", "capsula", "polvo",
    "sobre", "pieza", "piezas", "vitamina", "vitaminas", "acido", "sabor",
    "color", "calcetin", "parche", "balsamo", "barra", "protect", "dual",
    "generico", "infantil", "adulto", "spray", "jabon", "desodorante",
    "jeringa", "desechable", "insulina", "adhesivo", "pomada", "unguento",
    "micelar", "isotonica", "neutro", "repelente",
}

NO_CLASE_A = re.compile(
    r"\b(calcetin|calcetines|parche termico|unisex|repelente|gotas oftal|"
    r"solucion oftal|barra topica|balsamo en barra)\b",
    re.I,
)


def formas_compatibles(a: str | None, b: str | None) -> bool:
    if not a or not b:
        return True
    if a == b:
        return True
    return FORMA_FAMILIA.get(a) == FORMA_FAMILIA.get(b) and FORMA_FAMILIA.get(a) is not None


def norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s.lower()).strip()


def activos_sim(desc: str) -> list[str]:
    """Moléculas del genérico Similares: cabeza de cada segmento antes de la dosis."""
    cabeza = re.split(r"\d", desc or "", maxsplit=1)[0]
    ings = [sync.canonico(x) for x in sync.ingredientes(cabeza.replace("/", " + "))]
    toks = [sync.canonico(t) for t in sync.tokens_utiles(cabeza)]
    out: list[str] = []
    conocidos = INDISPENSABLE_TOKENS | CLASE_A
    for x in ings + toks:
        if x in PA_DEBIL or len(x) < 5:
            continue
        if x in ings or x in conocidos:
            if x not in out:
                out.append(x)
    return out


def clase_y_stock(sim: dict) -> tuple[str, int, int]:
    forma = (sim.get("forma") or "")
    jer = sim.get("jerarquia") or ""
    precio = float(sim.get("precio") or 0)
    desc = " ".join([sim.get("descripcion") or "", sim.get("marca") or ""])
    nd = norm(desc)

    linea_n = norm(sim.get("linea") or "")
    if NO_CLASE_A.search(desc):
        return "C", 2, 3
    if linea_n in {"higiene", "belleza"} or "shampoo" in nd or "desodorante" in nd:
        return "C", 2, 3
    if MARCA_PROPIA_SIMI.search(desc) or "suplement" in norm(jer):
        return "C", 2, 3
    if precio >= 250:
        return "C", 1, 2
    if precio >= 120:
        return "B", 2, 4
    if "inyect" in (forma or "") or "ampolleta" in nd or "ampolla" in nd:
        return "INY", 2, 3
    if (
        "curaci" in linea_n
        or "material de curaci" in linea_n
        or "material de curaci" in norm(jer)
    ):
        return "CUR", 6, 10
    if any(tok in nd for tok in CLASE_A):
        return "A", 6, 10
    linea = norm(sim.get("linea") or "")
    if any(x in linea for x in ("analges", "diabet", "cardio", "estomac", "gastro", "antihist")):
        return "A", 5, 8
    if any(x in linea for x in ("antibiot", "respirator", "antimicot")):
        return "B", 3, 6
    return "C", 2, 4


# Genéricos que sí se piden todos los días en mostrador (farmacia tipo Similares).
INDISPENSABLE_TOKENS = {
    "paracetamol", "ibuprofeno", "naproxeno", "diclofenaco", "ketorolaco", "metamizol",
    "acetilsalicilico",
    "omeprazol", "butilhioscina", "hioscina", "racecadotrilo",
    "metronidazol", "bismuto", "aluminio", "dimeticona", "mebendazol",
    "metformina", "glibenclamida",
    "losartan", "amlodipino", "enalapril", "captopril", "atenolol", "furosemida",
    "ambroxol", "dextrometorfano", "guaifenesina", "oxolamina", "clorfenamina",
    "loratadina", "cetirizina", "levocetirizina", "salbutamol", "fenilefrina",
    "amoxicilina", "ampicilina", "trimetoprima", "sulfametoxazol",
    "clotrimazol", "albendazol",
    "yodopovidona", "iodopovidona", "algodon", "termometro", "oxigenada",
}

EXCLUIR_SURTIR = re.compile(
    r"\b(6c|globulo|globulos|simi diab|simibaby|simiplaneta|simi chapulin|"
    r"cubreboca|tobillera|kitocream|kitocell|pirfen|candesartan|olmesartan|"
    r"nebivolol|lercanidipino|valsartan|pregabalina|gabapentina|mirtazapina|"
    r"risperidona|modafinilo|aripiprazol|desvenlafaxina|atomoxetina|"
    r"bicalutamida|mesalazina|orlistat|topiramato|oxcarbazepina|"
    r"hidroxocobalamina|norfenefrina|felodipino|ramipril|bisoprolol|"
    r"metoprolol|lisinopril|verapamilo|nifedipino|propranolol|clortalidona|"
    r"clort\b|isosorbida|metildopa|acarbosa|alendronato|itoprida|floroglucinol|"
    r"cuo protect|enteroger|sinuberase|bac claus|ultra fine|pluma|"
    r"oftal|otica|cabestrillo|kn95|nitrilo|genciana|pirfen|"
    r"sitagliptina|vildagliptina|linagliptina|dapagliflozina|empagliflozina|"
    r"ketoprofeno|sucralfato|parche|ampolleta|inyect)\b",
    re.I,
)

CURACION_INDISPENSABLE = re.compile(
    r"\b(yodopovidona|agua oxigenada|venda adhesiva|vendas adhesivas|"
    r"tela adhesiva|algodon|termometro digital|prueba de embarazo analog|"
    r"jeringa desechable 3ml|jeringa desechable 5ml)\b",
    re.I,
)


def por_que_surtir(desc: str) -> str:
    nd = norm(desc)
    pares = [
        (("paracetamol", "ibuprofeno", "naproxeno", "diclofenaco", "metamizol",
          "ketorolaco", "ketoprofeno", "acetilsalicilico"), "Analgésico de mostrador"),
        (("butilhioscina", "hioscina", "buscapina"), "Cólico / dolor de panza"),
        (("racecadotrilo",), "Diarrea — no hay hidrasec/genérico en anaquel"),
        (("metformina", "glibenclamida"), "Diabetes de mostrador"),
        (("ambroxol", "dextrometorfano", "guaifenesina", "oxolamina", "histiacil"),
         "Tos / gripa"),
        (("loratadina", "cetirizina", "clorfenamina", "levocetirizina"), "Alergia / gripa"),
        (("omeprazol", "sucralfato", "aluminio", "bismuto", "picot", "alka"),
         "Estómago / acidez"),
        (("amoxicilina", "ampicilina", "trimetoprima", "ceftriaxona", "metronidazol"),
         "Antibiótico de mostrador"),
        (("losartan", "amlodipino", "enalapril", "captopril", "atenolol"),
         "Presión / corazón de mostrador"),
        (("clotrimazol",), "Hongos / ginecológico"),
        (("mebendazol", "albendazol"), "Lombrices"),
        (("yodo", "oxigenada", "venda", "jeringa", "algodon", "termometro", "embarazo"),
         "Botiquín / curación que se acaba"),
    ]
    for toks, label in pares:
        if any(t in nd for t in toks):
            return label
    return "Alta rotación de mostrador"


def es_indispensable_mostrador(r: dict) -> bool:
    """True si el hueco es de los que hay que surtir ya (no especialidad)."""
    desc = r.get("descripcion") or ""
    nd = norm(desc)
    if EXCLUIR_SURTIR.search(desc) or EXCLUIR_SURTIR.search(nd):
        return False
    if r.get("clase") == "INY":
        return False
    # En mostrador el líquido de fiebre/dolor es ibuprofeno o Tempra, no diclofenaco.
    if "diclofenaco" in nd and re.search(r"\b(susp|jarabe|solucion)\b", nd):
        return False
    if MARCA_PROPIA_SIMI.search(desc) and "ambroxol" not in nd and "paracetamol" not in nd:
        return False
    if CURACION_INDISPENSABLE.search(desc) or CURACION_INDISPENSABLE.search(nd):
        return True
    toks = set(sync.tokens_utiles(desc))
    toks.update(r.get("ingredientes") or [])
    toks.update(activos_sim(desc))
    toks = {sync.canonico(t) for t in toks}
    return bool(toks & INDISPENSABLE_TOKENS)


BOTIQUIN_EQUIV = (
    (re.compile(r"jeringa desechable 3\s*ml|jeringa.{0,12}3\s*ml", re.I),
     re.compile(r"jeringa.{0,24}3\s*ml", re.I)),
    (re.compile(r"jeringa desechable 5\s*ml|jeringa.{0,12}5\s*ml", re.I),
     re.compile(r"jeringa.{0,24}5\s*ml", re.I)),
    (re.compile(r"agua oxigenada", re.I),
     re.compile(r"agua oxigenada|peroxido de hidrogeno", re.I)),
    (re.compile(r"\balgodon\b", re.I),
     re.compile(r"\balgodon\b", re.I)),
    (re.compile(r"tela adhesiva", re.I),
     re.compile(r"tela adhesiva", re.I)),
    (re.compile(r"venda adhesiva|vendas adhesivas", re.I),
     re.compile(r"curita|vendita|venda adhesiva|vendas adhesivas", re.I)),
    (re.compile(r"termometro", re.I),
     re.compile(r"termometro", re.I)),
    (re.compile(r"yodopovidona|iodopovidona", re.I),
     re.compile(r"yodopovidona|iodopovidona|isodine|dermodine", re.I)),
    (re.compile(r"prueba de embarazo", re.I),
     re.compile(r"prueba.{0,12}embarazo|meditest", re.I)),
)


def _blob_producto(p: dict) -> str:
    return norm(" ".join(filter(None, [
        str(p.get("nombre") or ""),
        str(p.get("principio_activo") or ""),
        str(p.get("forma_farmaceutica") or ""),
        str(p.get("concentracion") or ""),
        str(p.get("presentacion") or ""),
    ])))


def ya_hay_botiquin(desc: str, productos: list[dict]) -> bool:
    nd = norm(desc)
    for pat_sim, pat_fc in BOTIQUIN_EQUIV:
        if not pat_sim.search(nd):
            continue
        for p in productos:
            if int(p.get("stock") or 0) <= 0:
                continue
            if pat_fc.search(_blob_producto(p)):
                return True
    return False


def moleculas_texto(texto: str) -> list[str]:
    conocidos = INDISPENSABLE_TOKENS | CLASE_A
    ings = [sync.canonico(x) for x in sync.ingredientes(texto)]
    toks = [sync.canonico(t) for t in sync.tokens_utiles(texto)]
    out: list[str] = []
    for x in ings + toks:
        if x in PA_DEBIL or len(x) < 5:
            continue
        if x in ings or x in conocidos:
            if x not in out:
                out.append(x)
    return out


def ya_hay_equivalente(r: dict, idx_mol: dict[tuple[str, str], list[dict]],
                       productos: list[dict] | None = None) -> bool:
    """Si ya tenemos el mismo genérico en la misma familia de forma, no lo vuelvas a pedir.

    Excepción: pediátrico (suspensión/jarabe/gotas/supositorio) pide la misma vía.
    Excepción: metformina 1000 / 750 LP si en anaquel solo hay 500 u 850.
    """
    desc = r.get("descripcion") or ""
    if productos and ya_hay_botiquin(desc, productos):
        return True
    ings = [sync.canonico(x) for x in (r.get("ingredientes") or []) if x not in PA_DEBIL]
    ings = [x for x in ings if x]
    extra = activos_sim(desc)
    for x in extra:
        if x not in ings:
            ings.append(x)
    if not ings:
        return False
    forma = r.get("forma") or ""
    fam = FORMA_FAMILIA.get(forma) or forma or ""
    nd = norm(desc)
    pediatrico = forma in {"suspension", "jarabe", "gotas", "supositorio"} or bool(
        re.search(r"\b(ped|infantil|nino|bebe|gotas|supos)\b", nd)
    )
    candidatos = []
    for mol in ings:
        for (m, f), rows in idx_mol.items():
            if m != mol:
                continue
            if f == fam or (not fam) or (not f):
                candidatos.extend(rows)
    if not candidatos:
        return False
    if "metformina" in nd and re.search(r"\b(1000|1\s*gr|750)", nd):
        tienen = " ".join(
            norm(p.get("nombre") or "") + " " + str(p.get("concentracion") or "")
            for p in candidatos
        )
        if re.search(r"\b(1000|1\s*gr)", nd):
            return bool(re.search(r"\b(1000|1\s*g\b|1\s*gr)", tienen))
        if re.search(r"\b750", nd):
            return "750" in tienen
    if pediatrico:
        forma_ok = {forma} if forma else set()
        if forma in {"suspension", "jarabe"}:
            forma_ok.update({"suspension", "jarabe"})
        if forma == "gotas":
            forma_ok.add("gotas")
        if forma == "supositorio":
            forma_ok.add("supositorio")
        return any(
            sync.extraer_forma((p.get("nombre") or "") + " " + (p.get("forma_farmaceutica") or ""))
            in forma_ok
            for p in candidatos
            if int(p.get("stock") or 0) > 0
        )
    return any(int(p.get("stock") or 0) > 0 for p in candidatos)


def indice_molecula(productos: list[dict]) -> dict[tuple[str, str], list[dict]]:
    idx: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for p in productos:
        texto = " ".join(filter(None, [
            str(p.get("principio_activo") or ""),
            str(p.get("nombre") or ""),
            str(p.get("forma_farmaceutica") or ""),
        ]))
        ings = moleculas_texto(texto)
        forma = sync.extraer_forma(texto) or sync.extraer_forma(p.get("forma_farmaceutica") or "")
        fam = FORMA_FAMILIA.get(forma or "") or (forma or "")
        for ing in ings:
            idx[(ing, fam)].append(p)
    return idx


def elegir_pedido_surtir(huecos: list[dict], productos: list[dict]) -> list[dict]:
    """Faltantes que sí hay que comprar ahora: alta rotación y sin equivalente en anaquel."""
    idx = indice_molecula(productos)
    out = []
    vistos: set[str] = set()
    for r in huecos:
        if not es_indispensable_mostrador(r):
            continue
        if ya_hay_equivalente(r, idx, productos):
            continue
        ings = activos_sim(r.get("descripcion") or "") or (r.get("ingredientes") or [])
        # Un renglón por genérico+forma+dosis redondeada, no 3 SKUs Simi del mismo.
        clave = (
            " ".join(sorted(ings))
            + "|"
            + (r.get("forma") or "")
            + "|"
            + " ".join(sorted(r.get("concentraciones") or [])[:3])
        )
        if clave in vistos and "jeringa" not in norm(r.get("descripcion") or ""):
            continue
        vistos.add(clave)
        rec = dict(r)
        rec["por_que"] = por_que_surtir(r.get("descripcion") or "")
        rec["pedir"] = 6 if r.get("clase") in ("A", "CUR") else max(int(r.get("pedir") or 3), 3)
        if "metformina" in norm(rec["descripcion"]) and re.search(r"1000|1\s*gr", norm(rec["descripcion"])):
            rec["pedir"] = 10
        out.append(rec)
    out.sort(key=lambda x: (x.get("por_que") or "", -(x.get("precio") or 0)))
    return out


def _vtex_get(url: str) -> tuple[list, str]:
    req = urllib.request.Request(
        url,
        headers={"Accept": "application/json", "User-Agent": "Mozilla/5.0 (FarmaCapital cruce)"},
    )
    with urllib.request.urlopen(req, timeout=40) as resp:
        resources = resp.headers.get("resources") or resp.headers.get("Resources") or ""
        return json.loads(resp.read()), resources


def cargar_similares_vtex() -> list[dict]:
    """Catálogo público VTEX, paginado de 50 en 50."""
    out: dict[str, dict] = {}
    start = 0
    page = 50
    total = None
    while True:
        end = start + page - 1
        url = f"{VTEX_SEARCH}?_from={start}&_to={end}"
        items, resources = _vtex_get(url)
        if total is None and "/" in resources:
            try:
                total = int(resources.split("/")[-1])
            except ValueError:
                total = None
        if not items:
            break
        for it in items:
            rec = vtex_a_fila(it)
            if rec:
                out[rec["sim_sku"]] = rec
        print(f"  VTEX {start}-{start + len(items) - 1}" + (f" / {total}" if total else ""), flush=True)
        start += page
        if total is not None and start >= total:
            break
        if len(items) < page:
            break
        time.sleep(0.15)
    return list(out.values())


def vtex_a_fila(it: dict) -> dict | None:
    nombre = (it.get("productName") or "").strip()
    if not nombre:
        return None
    cats = it.get("categories") or []
    cat_txt = " ".join(cats)
    if SKIP_CAT_RE.search(cat_txt):
        return None
    jerarquia = ""
    linea = ""
    if cats:
        partes = [p for p in cats[0].split("/") if p]
        jerarquia = partes[0] if partes else ""
        linea = partes[1] if len(partes) > 1 else partes[0] if partes else ""
    items = it.get("items") or [{}]
    sellers = (items[0].get("sellers") or [])
    precio = 0.0
    if sellers:
        precio = float(sellers[0].get("commertialOffer", {}).get("Price") or 0)
    sku = str(it.get("productReference") or it.get("productId") or "").strip()
    marca = (it.get("brand") or "").strip()
    if norm(marca) in MARCAS_CASA_SIMI:
        marca = ""
    return {
        "sim_sku": sku,
        "fuente": "vtex",
        "marca": marca,
        "marca_norm": sync.normalizar(marca),
        "descripcion": nombre,
        "desc_norm": sync.normalizar(nombre),
        "precio": precio,
        "concentracion": "",
        "contenido": "",
        "presentacion": (items[0].get("name") or ""),
        "linea": linea,
        "jerarquia": jerarquia,
        "grupo": "",
        "concentraciones": sync.extraer_concentraciones(nombre),
        "cantidad": sync.extraer_cantidad(nombre),
        "forma": sync.extraer_forma(nombre),
        "ingredientes": activos_sim(nombre),
        "n_marcas": 1,
        "link": it.get("link") or "",
        "requiere_receta": ",".join(it.get("Requiere Receta") or []),
    }


def cargar_similares_excel(ruta: Path) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    by_sku: dict[str, dict] = {}
    for fila in ws.iter_rows(min_row=3, values_only=True):
        sku = str(fila[0] or "").strip()
        if not sku or sku == "None":
            continue
        estatus = str(fila[13] or "").strip().upper() if len(fila) > 13 else "ACTIVO"
        if estatus and estatus not in ("ACTIVO", ""):
            continue
        marca = (fila[1] or "").strip() if isinstance(fila[1], str) else ""
        desc = (fila[2] or "").strip() if isinstance(fila[2], str) else ""
        if not desc:
            continue
        try:
            precio = float(str(fila[3]).replace("$", "").replace(",", ""))
        except (TypeError, ValueError):
            precio = 0.0
        concentracion = str(fila[4] or "").strip()
        contenido = str(fila[5] or "").strip()
        presentacion = str(fila[7] or "").strip() if len(fila) > 7 else ""
        linea = str(fila[9] or "").strip() if len(fila) > 9 else ""
        jerarquia = str(fila[10] or "").strip() if len(fila) > 10 else ""
        grupo = str(fila[11] or "").strip() if len(fila) > 11 else ""
        if SKIP_CAT_RE.search(f"{jerarquia} {linea}"):
            continue
        if jerarquia and not INCLUDE_CAT_RE.search(jerarquia) and jerarquia.upper() not in {
            "MEDICAMENTOS", "MATERIAL DE CURACION", "SUPLEMENTOS", "SALUD SEXUAL",
            "MATERNIDAD", "REMEDIOS HERBOLARIOS",
        }:
            continue
        cand = {
            "sim_sku": sku,
            "fuente": "excel",
            "marca": marca,
            "marca_norm": sync.normalizar(marca),
            "descripcion": desc,
            "desc_norm": sync.normalizar(desc),
            "precio": precio,
            "concentracion": concentracion,
            "contenido": contenido,
            "presentacion": presentacion,
            "linea": linea,
            "jerarquia": jerarquia,
            "grupo": grupo,
            "concentraciones": sync.extraer_concentraciones(f"{concentracion} {desc}"),
            "cantidad": sync.extraer_cantidad(contenido) or sync.extraer_cantidad(desc),
            "forma": sync.extraer_forma(presentacion) or sync.extraer_forma(desc),
            "ingredientes": activos_sim(desc),
            "n_marcas": 1,
            "link": "",
            "requiere_receta": "",
        }
        prev = by_sku.get(sku)
        if prev is None:
            by_sku[sku] = cand
        else:
            prev["n_marcas"] += 1
            if len(desc) > len(prev["descripcion"]):
                cand["n_marcas"] = prev["n_marcas"]
                by_sku[sku] = cand
    wb.close()
    return list(by_sku.values())


def fetch_productos_stock(url: str, key: str) -> list[dict]:
    H = {
        "apikey": key,
        "Authorization": f"Bearer {key}",
        "Accept": "application/json",
        "Prefer": "count=exact",
    }
    select = (
        "id,sku,nombre,principio_activo,marca,presentacion,"
        "forma_farmaceutica,concentracion,costo,precio,categoria,tipo,stock,stock_minimo,activo"
    )
    page = 250
    out = []
    start = 0
    while True:
        last_err: Exception | None = None
        chunk: list = []
        for intento in range(6):
            try:
                req = urllib.request.Request(
                    url + f"/rest/v1/productos?select={select}"
                    "&activo=eq.true&order=nombre.asc",
                    headers={**H, "Range": f"{start}-{start + page - 1}"},
                )
                with urllib.request.urlopen(req, timeout=45) as r:
                    chunk = json.loads(r.read())
                last_err = None
                break
            except Exception as exc:
                last_err = exc
                time.sleep(2 * (intento + 1))
        if last_err:
            raise last_err
        if not isinstance(chunk, list):
            raise RuntimeError("Respuesta inesperada de productos")
        out.extend(chunk)
        if len(chunk) < page:
            break
        start += page
    return out


def token_index(productos: list[dict], atts: list[dict]) -> dict[str, list[int]]:
    idx: dict[str, list[int]] = defaultdict(list)
    for i, (_p, att) in enumerate(zip(productos, atts)):
        toks = set(sync.tokens_utiles(att["texto"]))
        toks.update(att["ingredientes"])
        for t in toks:
            if len(t) >= 4:
                idx[t].append(i)
    return idx


def nuestros_atts(productos: list[dict]) -> list[dict]:
    atts = []
    for p in productos:
        att = ex.nuestros_atributos(p)
        extra = " ".join(filter(None, [
            str(p.get("concentracion") or ""),
            str(p.get("forma_farmaceutica") or ""),
        ]))
        if extra:
            att["texto"] = (att["texto"] + " " + extra).strip()
            att["concentraciones"] |= sync.extraer_concentraciones(extra)
            att["cantidad"] = att["cantidad"] or sync.extraer_cantidad(att["texto"])
            att["forma"] = att["forma"] or sync.extraer_forma(extra)
        atts.append(att)
    return atts


def cruzar(sim: list[dict], productos: list[dict]) -> tuple[list[dict], list[dict], list[dict]]:
    freq: Counter[str] = Counter()
    for c in sim:
        freq.update(set(sync.tokens_utiles(c["descripcion"])))
    minimo = max(25, int(len(sim) * 0.02))
    vagos = {t for t, n in freq.items() if n >= minimo}

    atts = nuestros_atts(productos)
    idx = token_index(productos, atts)

    covered: dict[str, dict] = {}
    for s in sim:
        toks = set(sync.tokens_utiles(s["descripcion"]))
        toks.update(s.get("ingredientes") or [])
        cand_i: set[int] = set()
        for t in toks:
            if len(t) >= 4:
                cand_i.update(idx.get(t, []))
        best = None
        best_score = 0
        best_raz: list[str] = []
        best_equiv = False
        for i in cand_i:
            score, razones, _ = ex.evaluar(atts[i], s, vagos)
            sim_ing = [
                sync.canonico(x)
                for x in (s.get("ingredientes") or [])
                if x not in vagos and x not in PA_DEBIL and len(x) >= 5
            ]
            fc_ing = [
                sync.canonico(x) for x in atts[i]["ingredientes"]
                if x not in PA_DEBIL and len(x) >= 5
            ]
            fc_txt = " ".join(atts[i]["ingredientes"] + sync.tokens_utiles(atts[i]["texto"]))
            faltan = []
            if sim_ing:
                faltan = [x for x in sim_ing if x not in fc_ing and x not in fc_txt]
                if faltan:
                    score = min(score, 55)
                    razones = razones + [f"faltan en tu ficha: {', '.join(faltan)}"]
                if len(sim_ing) >= 2 and len(fc_ing) == 1:
                    score = min(score, 58)
                    razones = razones + ["combinado Similares vs monofármaco tuyo"]
            pa_ok = bool(sim_ing) and not faltan
            conc_ok = bool(atts[i]["concentraciones"] & s["concentraciones"]) if (
                atts[i]["concentraciones"] and s["concentraciones"]
            ) else False
            forma_ok = formas_compatibles(atts[i]["forma"], s["forma"])
            fc_nom = sync.normalizar(productos[i].get("nombre") or "")
            nombre_set = ex.fuzz.token_set_ratio(s["desc_norm"], fc_nom)
            nombre_ratio = ex.fuzz.ratio(s["desc_norm"], fc_nom)
            # token_set_ratio("tempra", "paracetamol 80mg ... tempra") = 100
            # y un SKU vago no puede tapar 4 presentaciones distintas.
            nombre_corto = len(fc_nom.split()) <= 2
            if nombre_corto and not conc_ok:
                pass
            elif nombre_set >= 88 and nombre_ratio >= 55:
                score = max(score, 82)
                razones = razones + [f"nombre muy parecido ({nombre_set:.0f}/{nombre_ratio:.0f})"]
            elif nombre_set >= 80 and nombre_ratio >= 48 and (conc_ok or not s["concentraciones"]):
                score = max(score, 76)
                razones = razones + [f"nombre parecido ({nombre_set:.0f}/{nombre_ratio:.0f})"]
            # Jeringas / material: el volumen tiene que coincidir, no basta la palabra.
            nd = norm(s["descripcion"])
            if "jeringa" in nd and s["concentraciones"] and not conc_ok:
                score = min(score, 50)
                razones = razones + ["volumen de jeringa no coincide"]
            equiv = pa_ok and conc_ok and forma_ok and bool(sim_ing)
            if equiv and score < UMBRAL_EQUIV:
                score = max(score, UMBRAL_EQUIV)
                razones = razones + ["equivalente PA+concentración+forma"]
            if score == 0 or (razones == ["sin identidad comprobable"] and nombre_set < 80):
                continue
            if score > best_score:
                best_score = score
                best = productos[i]
                best_raz = razones
                best_equiv = equiv
        if best is not None and (best_score >= UMBRAL_CUBIERTO or best_equiv):
            covered[s["sim_sku"]] = {
                "score": best_score,
                "razones": "; ".join(best_raz),
                "prod": best,
                "equiv": best_equiv,
            }

    huecos, rellenar, ok = [], [], []
    for s in sim:
        clase, minimo_s, objetivo = clase_y_stock(s)
        rec = {**s, "clase": clase, "minimo": minimo_s, "objetivo": objetivo}
        hit = covered.get(s["sim_sku"])
        if not hit:
            rec.update({
                "estado": "HUECO",
                "fc_sku": "",
                "fc_nombre": "",
                "fc_marca": "",
                "fc_pa": "",
                "stock": 0,
                "pedir": minimo_s,
                "match": 0,
                "razones": "",
            })
            huecos.append(rec)
            continue
        p = hit["prod"]
        stock = int(p.get("stock") or 0)
        rec.update({
            "fc_sku": p.get("sku") or "",
            "fc_nombre": p.get("nombre") or "",
            "fc_marca": p.get("marca") or "",
            "fc_pa": p.get("principio_activo") or "",
            "stock": stock,
            "match": hit["score"],
            "razones": hit["razones"],
        })
        if stock < minimo_s:
            rec["estado"] = "RELLENAR"
            rec["pedir"] = minimo_s - stock
            rellenar.append(rec)
        else:
            rec["estado"] = "OK"
            rec["pedir"] = 0
            ok.append(rec)
    return huecos, rellenar, ok


def prio_key(r: dict):
    order = {"A": 0, "INY": 1, "CUR": 2, "B": 3, "C": 4}
    return (order.get(r["clase"], 9), -(r.get("precio") or 0), r["descripcion"])


def escribir_csv(path: Path, rows: list[dict], extra_fc: bool) -> None:
    headers = [
        "prioridad", "estado", "clase", "pedir", "sku_similares", "generico_similares",
        "concentraciones", "piezas", "forma", "linea", "jerarquia",
        "precio_similares", "min_sucursal", "objetivo",
    ]
    if extra_fc:
        headers += ["fc_sku", "fc_nombre", "fc_marca", "stock", "match", "razones"]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for i, r in enumerate(rows, start=1):
            conc = " ".join(sorted(r.get("concentraciones") or []))
            vals = [
                i, r.get("estado"), r["clase"], r.get("pedir") or 0, r["sim_sku"],
                r["descripcion"], conc, r.get("cantidad") or "", r.get("forma") or "",
                r.get("linea") or "", r.get("jerarquia") or "",
                r.get("precio") or 0, r["minimo"], r["objetivo"],
            ]
            if extra_fc:
                vals += [
                    r.get("fc_sku") or "", r.get("fc_nombre") or "",
                    r.get("fc_marca") or "", r.get("stock") or 0,
                    r.get("match") or 0, r.get("razones") or "",
                ]
            w.writerow(vals)


def escribir_snapshot(path: Path, sim: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=[
            "sim_sku", "descripcion", "marca", "precio", "linea", "jerarquia",
            "forma", "cantidad", "concentraciones", "ingredientes", "link", "fuente",
        ])
        w.writeheader()
        for r in sim:
            w.writerow({
                "sim_sku": r["sim_sku"],
                "descripcion": r["descripcion"],
                "marca": r.get("marca") or "",
                "precio": r.get("precio") or 0,
                "linea": r.get("linea") or "",
                "jerarquia": r.get("jerarquia") or "",
                "forma": r.get("forma") or "",
                "cantidad": r.get("cantidad") or "",
                "concentraciones": " ".join(sorted(r.get("concentraciones") or [])),
                "ingredientes": " + ".join(r.get("ingredientes") or []),
                "link": r.get("link") or "",
                "fuente": r.get("fuente") or "",
            })


def escribir_md(path: Path, resumen: dict, huecos: list[dict]) -> None:
    por_linea: Counter[str] = Counter()
    for r in huecos:
        if r["clase"] in ("A", "B", "CUR"):
            por_linea[r.get("linea") or r.get("jerarquia") or "Sin línea"] += 1
    lineas = [
        f"# Cruce FarmaCapital vs surtido Similares ({resumen['fecha']})",
        "",
        f"Fuente: **{resumen['fuente']}**.",
        "",
        "El match es por **genérico** (principio + concentración + forma), no por marca comercial.",
        "Si Similares vende ibuprofeno 400 mg 10 tabletas y nosotros tenemos AMSA/Ultra de esa misma presentación, cuenta como cubierto.",
        "",
        "## Resumen",
        "",
        f"- Artículos únicos Similares (farmacia, sin souvenirs/perfumería/alimentos): **{resumen['similares_unicos']}**",
        f"- Productos activos en nuestro inventario: **{resumen['inventario_activo']}**",
        f"- Ya cubiertos (los tenemos, aunque sea otra marca): **{resumen['cubiertos']}** ({resumen['cobertura_pct']}%)",
        f"  - con stock suficiente: **{resumen['ok']}**",
        f"  - hay que rellenar (bajo el mínimo sucursal): **{resumen['rellenar']}**",
        f"- Huecos (Similares lo vende y nosotros no): **{resumen['huecos']}**",
        f"  - clase A (rotación alta): **{resumen['huecos_a']}**",
        f"  - clase B: **{resumen['huecos_b']}**",
        f"  - clase C: **{resumen['huecos_c']}**",
        f"  - inyectables: **{resumen['huecos_iny']}**",
        f"  - curación: **{resumen['huecos_cur']}**",
        "",
        "Para igualar el surtido de mostrador de Similares, el pedido útil es **clase A + B + curación**",
        f"({resumen['huecos_a'] + resumen['huecos_b'] + resumen['huecos_cur']} huecos).",
        "Clase C (especialidad / marca propia Simi) se pide de 2 en 2 cuando ya hay consulta.",
        "",
        "Archivos:",
        "",
        f"- `{resumen.get('xlsx_surtir') or ''}` — **Excel para surtir** (solo alta rotación de mostrador)",
        f"- `{resumen.get('xlsx_pedido') or ''}` — cruce completo (A/B/curación + rellenar)",
        f"- `{resumen['csv_prioridad']}` — huecos A + B + curación (comprar primero)",
        f"- `{resumen['csv_huecos']}` — resto de huecos (especialidad)",
        f"- `{resumen['csv_rellenar']}` — ya los tenemos, stock bajo el mínimo",
        f"- `{resumen['csv_cubiertos']}` — ya cubiertos",
        "",
        "## Huecos prioridad por línea",
        "",
        "| Línea | Huecos A/B/curación |",
        "| --- | ---: |",
    ]
    for linea, n in por_linea.most_common(20):
        lineas.append(f"| {linea} | {n} |")
    lineas += [
        "",
        "## Huecos clase A (muestra)",
        "",
        "| Pedir | Precio Simi | Genérico | Línea |",
        "| ---: | ---: | --- | --- |",
    ]
    n = 0
    for r in huecos:
        if r["clase"] != "A":
            continue
        lineas.append(
            f"| {r['pedir']} | ${r['precio']:.0f} | {r['descripcion'][:70]} | {r.get('linea') or ''} |"
        )
        n += 1
        if n >= 40:
            break
    if n == 0:
        lineas.append("| — | — | (no hubo huecos clase A) | — |")
    path.write_text("\n".join(lineas) + "\n", encoding="utf-8")


def escribir_xlsx_pedido(path: Path, resumen: dict, huecos: list[dict], rellenar: list[dict]) -> None:
    """Excel para surtir: alta rotación primero, luego media, luego rellenar."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    fill_h = PatternFill("solid", fgColor="1A1A1A")
    font_h = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    font_t = Font(name="Calibri", bold=True, size=16)
    thin = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )
    fill_a = PatternFill("solid", fgColor="F8D7DA")
    fill_b = PatternFill("solid", fgColor="FFF3CD")
    fill_cur = PatternFill("solid", fgColor="D6EAF8")

    def style_header(ws, ncols):
        for col in range(1, ncols + 1):
            cell = ws.cell(1, col)
            cell.fill = fill_h
            cell.font = font_h
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}1"
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 22

    def fill_clase(clase):
        return {"A": fill_a, "B": fill_b, "CUR": fill_cur, "INY": fill_cur}.get(clase)

    alta = [r for r in huecos if r["clase"] == "A"]
    media = [r for r in huecos if r["clase"] in ("B", "CUR")]
    rell_ab = [r for r in rellenar if r["clase"] in ("A", "B", "CUR")]

    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = "Pedido para igualar surtido Similares (alta rotación)"
    ws["A1"].font = font_t
    ws.merge_cells("A1:B1")
    ws["A2"] = f"Fuente: {resumen['fuente']} · inventario vivo {resumen['fecha']}"
    ws["A3"] = "Compara por genérico (principio + concentración + forma), no por marca comercial."
    rows_r = [
        ("Genéricos únicos Similares (farmacia)", resumen["similares_unicos"]),
        ("Productos activos en tu inventario", resumen["inventario_activo"]),
        ("Ya cubiertos", resumen["cubiertos"]),
        ("Cobertura", f"{resumen['cobertura_pct']}%"),
        ("Huecos totales", resumen["huecos"]),
        ("  · clase A — comprar primero (alta rotación)", resumen["huecos_a"]),
        ("  · clase B — rotación media", resumen["huecos_b"]),
        ("  · curación", resumen["huecos_cur"]),
        ("Piezas a pedir — solo clase A", sum(r["pedir"] for r in alta)),
        ("Piezas a pedir — A + B + curación", sum(r["pedir"] for r in alta + media)),
        ("Piezas a rellenar (ya los tienes, stock bajo)", sum(r["pedir"] for r in rell_ab)),
    ]
    ws["A5"] = "Métrica"
    ws["B5"] = "Valor"
    style_header(ws, 2)
    for i, (a, b) in enumerate(rows_r, start=6):
        ws.cell(i, 1, a)
        ws.cell(i, 2, b)
    ws["A19"] = (
        "Hoja 1_Alta_rotacion: los indispensables de mostrador que Similares sí tiene y tú no. "
        "Pide esas piezas (columna Pedir) a Levic / Farmalive / AMSA del mismo genérico. "
        "No copies la marca propia Simi: busca el equivalente genérico."
    )
    ws.merge_cells("A19:E21")
    ws["A19"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 64
    ws.column_dimensions["B"].width = 18

    def sheet_pedido(name, data, extra_fc=False):
        ws = wb.create_sheet(name)
        headers = [
            "Prioridad", "Clase", "Pedir", "SKU Similares", "Genérico",
            "Concentración", "Contenido", "Forma", "Línea",
            "Precio venta Similares",
        ]
        if extra_fc:
            headers += ["Tu SKU", "Tu producto", "Stock hoy"]
        for i, h in enumerate(headers, 1):
            ws.cell(1, i, h)
        style_header(ws, len(headers))
        for r_i, r in enumerate(data, start=2):
            conc = r.get("concentracion") or " ".join(sorted(r.get("concentraciones") or []))
            contenido = r.get("contenido") or r.get("presentacion") or ""
            vals = [
                r_i - 1,
                r["clase"],
                r["pedir"],
                r["sim_sku"],
                r["descripcion"],
                conc,
                contenido,
                r.get("forma") or "",
                r.get("linea") or "",
                r.get("precio") or 0,
            ]
            if extra_fc:
                vals += [r.get("fc_sku") or "", r.get("fc_nombre") or "", r.get("stock") or 0]
            for c_i, v in enumerate(vals, 1):
                cell = ws.cell(r_i, c_i, v)
                cell.border = thin
                if c_i == 2:
                    fill = fill_clase(r["clase"])
                    if fill:
                        cell.fill = fill
            ws.cell(r_i, 3).font = Font(name="Calibri", bold=True, size=12)
        widths = [10, 8, 8, 14, 46, 18, 22, 12, 24, 16, 16, 36, 12]
        for i, w in enumerate(widths[:len(headers)], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        if data:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(data)+1}"
        return ws

    sheet_pedido("1_Alta_rotacion", alta)
    sheet_pedido("2_Media_y_curacion", media)
    sheet_pedido("3_Rellenar_lo_que_ya_tienes", rell_ab, extra_fc=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def escribir_xlsx_surtir(path: Path, resumen: dict, pedido: list[dict]) -> None:
    """Excel corto para surtir: solo indispensables de alta rotación."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    fill_h = PatternFill("solid", fgColor="1A1A1A")
    font_h = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    thin = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )
    fill_alt = PatternFill("solid", fgColor="FFF8E7")

    ws = wb.active
    ws.title = "COMPRAR"
    headers = [
        "Pedir", "Qué comprar (genérico)", "Concentración / contenido",
        "Forma", "Línea", "Para qué", "Precio ref. Similares", "SKU Similares",
    ]
    for i, h in enumerate(headers, 1):
        c = ws.cell(1, i, h)
        c.fill = fill_h
        c.font = font_h
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 24
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    for r_i, r in enumerate(pedido, start=2):
        conc = (r.get("concentracion") or "").strip()
        if not conc:
            conc = " ".join(sorted(r.get("concentraciones") or []))
        contenido = r.get("contenido") or r.get("presentacion") or ""
        pres = " · ".join(x for x in [conc, contenido] if x)
        vals = [
            r["pedir"],
            r["descripcion"],
            pres,
            r.get("forma") or "",
            r.get("linea") or "",
            r.get("por_que") or "",
            r.get("precio") or 0,
            r.get("sim_sku") or "",
        ]
        for c_i, v in enumerate(vals, 1):
            cell = ws.cell(r_i, c_i, v)
            cell.border = thin
            if r_i % 2 == 0:
                cell.fill = fill_alt
        ws.cell(r_i, 1).font = Font(name="Calibri", bold=True, size=14)

    widths = [8, 48, 28, 12, 22, 36, 16, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    if pedido:
        ws.auto_filter.ref = f"A1:H{len(pedido)+1}"

    res = wb.create_sheet("Resumen", 0)
    res["A1"] = "Pedido de alta rotación — FarmaCapital vs Similares"
    res["A1"].font = Font(name="Calibri", bold=True, size=16)
    res.merge_cells("A1:B1")
    res["A2"] = (
        "Solo lo que se pide en mostrador y hoy no tenemos (ni un genérico equivalente). "
        "Compra el mismo principio/dosis/forma en Levic, Farmalive o AMSA. "
        "No copies la marca propia Simi."
    )
    res.merge_cells("A2:B4")
    res["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    filas = [
        ("Fecha", resumen["fecha"]),
        ("Fuente", resumen["fuente"]),
        ("Renglones a comprar", len(pedido)),
        ("Piezas a pedir", sum(r["pedir"] for r in pedido)),
        ("Huecos totales Similares (referencia)", resumen["huecos"]),
        ("De esos, alta rotación a surtir", len(pedido)),
    ]
    res["A6"] = "Métrica"
    res["B6"] = "Valor"
    for i in (6,):
        for col in range(1, 3):
            res.cell(i, col).fill = fill_h
            res.cell(i, col).font = font_h
    for i, (a, b) in enumerate(filas, start=7):
        res.cell(i, 1, a)
        res.cell(i, 2, b)
    por = Counter(r.get("por_que") or "" for r in pedido)
    res["A15"] = "Por rubro"
    res["A15"].font = Font(name="Calibri", bold=True, size=12)
    res["A16"] = "Rubro"
    res["B16"] = "Renglones"
    for i, h in enumerate(["A16", "B16"], 1):
        pass
    for col in range(1, 3):
        res.cell(16, col).fill = fill_h
        res.cell(16, col).font = font_h
    for i, (k, n) in enumerate(por.most_common(), start=17):
        res.cell(i, 1, k)
        res.cell(i, 2, n)
    res.column_dimensions["A"].width = 56
    res.column_dimensions["B"].width = 22

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def cargar_snapshot_csv(ruta: Path) -> list[dict]:
    out = []
    with ruta.open(encoding="utf-8") as f:
        for row in csv.DictReader(f):
            desc = row.get("descripcion") or ""
            concs = set((row.get("concentraciones") or "").split())
            ings = activos_sim(desc)
            out.append({
                "sim_sku": row.get("sim_sku") or "",
                "fuente": row.get("fuente") or "snapshot",
                "marca": row.get("marca") or "",
                "marca_norm": sync.normalizar(row.get("marca") or ""),
                "descripcion": desc,
                "desc_norm": sync.normalizar(desc),
                "precio": float(row.get("precio") or 0),
                "concentracion": "",
                "contenido": "",
                "presentacion": "",
                "linea": row.get("linea") or "",
                "jerarquia": row.get("jerarquia") or "",
                "grupo": "",
                "concentraciones": concs or sync.extraer_concentraciones(desc),
                "cantidad": int(row["cantidad"]) if str(row.get("cantidad") or "").isdigit() else sync.extraer_cantidad(desc),
                "forma": row.get("forma") or sync.extraer_forma(desc),
                "ingredientes": ings or activos_sim(desc),
                "n_marcas": 1,
                "link": row.get("link") or "",
                "requiere_receta": "",
            })
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--excel", default="", help="Ruta al Excel de artículos (opcional)")
    ap.add_argument("--snapshot", default="", help="CSV VTEX ya bajado (salta la API)")
    args = ap.parse_args()

    env = sync.cargar_env()
    url = env.get("REACT_APP_SUPABASE_URL", "").rstrip("/")
    key = env.get("SUPABASE_SERVICE_ROLE_KEY") or env.get("REACT_APP_SUPABASE_ANON_KEY", "")
    if not url or not key:
        sys.exit("Faltan credenciales Supabase en .env")

    excel = Path(args.excel).expanduser() if args.excel else EXCEL_DEFAULT
    snap = Path(args.snapshot).expanduser() if args.snapshot else Path()
    if excel.exists():
        print(f"Usando Excel {excel}")
        sim = cargar_similares_excel(excel)
        fuente = f"excel:{excel.name}"
    elif snap.exists():
        print(f"Usando snapshot {snap}")
        sim = cargar_snapshot_csv(snap)
        fuente = f"vtex:farmaciasdesimilares.com ({snap.name})"
    else:
        print("Excel no está en disco (gitignored). Bajando catálogo VTEX de Similares…")
        sim = cargar_similares_vtex()
        fuente = "vtex:farmaciasdesimilares.com"

    print(f"Similares únicos: {len(sim)}")
    productos = fetch_productos_stock(url, key)
    print(f"Inventario activo: {len(productos)}")

    huecos, rellenar, ok = cruzar(sim, productos)
    huecos.sort(key=prio_key)
    rellenar.sort(key=prio_key)
    ok.sort(key=lambda r: r["descripcion"])

    tag = HOY.replace("-", "")
    csv_prioridad = f"pricing/reportes/cruce_similares_{tag}_prioridad.csv"
    csv_huecos = f"pricing/reportes/cruce_similares_{tag}_huecos.csv"
    csv_rellenar = f"pricing/reportes/cruce_similares_{tag}_rellenar.csv"
    csv_cubiertos = f"pricing/reportes/cruce_similares_{tag}_cubiertos.csv"
    md_path = REPORTE_DIR / f"cruce_similares_{tag}.md"
    snap_path = SNAPSHOT_DIR / f"similares_catalogo_{tag}.csv"

    p1 = [r for r in huecos if r["clase"] in ("A", "B", "CUR")]
    p_resto = [r for r in huecos if r["clase"] not in ("A", "B", "CUR")]

    xlsx_pedido = REPORTE_DIR / f"cruce_similares_pedido_{tag}.xlsx"
    escribir_csv(ROOT / csv_prioridad, p1, extra_fc=False)
    escribir_csv(ROOT / csv_huecos, p_resto, extra_fc=False)
    escribir_csv(ROOT / csv_rellenar, rellenar, extra_fc=True)
    escribir_csv(ROOT / csv_cubiertos, ok, extra_fc=True)
    if fuente.startswith("excel:"):
        escribir_snapshot(SNAPSHOT_DIR / f"similares_excel_{tag}.csv", sim)
    else:
        escribir_snapshot(snap_path, sim)

    resumen = {
        "fecha": HOY,
        "fuente": fuente,
        "similares_unicos": len(sim),
        "inventario_activo": len(productos),
        "cubiertos": len(ok) + len(rellenar),
        "ok": len(ok),
        "rellenar": len(rellenar),
        "huecos": len(huecos),
        "huecos_a": sum(1 for r in huecos if r["clase"] == "A"),
        "huecos_b": sum(1 for r in huecos if r["clase"] == "B"),
        "huecos_c": sum(1 for r in huecos if r["clase"] == "C"),
        "huecos_iny": sum(1 for r in huecos if r["clase"] == "INY"),
        "huecos_cur": sum(1 for r in huecos if r["clase"] == "CUR"),
        "cobertura_pct": round(100 * (len(ok) + len(rellenar)) / len(sim), 1) if sim else 0,
        "csv_prioridad": csv_prioridad,
        "csv_huecos": csv_huecos,
        "csv_rellenar": csv_rellenar,
        "csv_cubiertos": csv_cubiertos,
        "xlsx_pedido": str(xlsx_pedido.relative_to(ROOT)),
    }
    pedido = elegir_pedido_surtir(huecos, productos)
    xlsx_surtir = REPORTE_DIR / f"pedido_alta_rotacion_surtir_{tag}.xlsx"
    csv_surtir = REPORTE_DIR / f"pedido_alta_rotacion_surtir_{tag}.csv"
    escribir_xlsx_pedido(xlsx_pedido, resumen, huecos, rellenar)
    escribir_xlsx_surtir(xlsx_surtir, resumen, pedido)
    with csv_surtir.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow([
            "pedir", "generico", "concentracion_contenido", "forma", "linea",
            "para_que", "precio_similares", "sku_similares",
        ])
        for r in pedido:
            conc = r.get("concentracion") or " ".join(sorted(r.get("concentraciones") or []))
            contenido = r.get("contenido") or r.get("presentacion") or ""
            w.writerow([
                r["pedir"], r["descripcion"], " · ".join(x for x in [conc, contenido] if x),
                r.get("forma") or "", r.get("linea") or "", r.get("por_que") or "",
                r.get("precio") or 0, r.get("sim_sku") or "",
            ])
    resumen["xlsx_surtir"] = str(xlsx_surtir.relative_to(ROOT))
    resumen["pedido_surtir"] = len(pedido)
    resumen["pzas_surtir"] = sum(r["pedir"] for r in pedido)
    escribir_md(md_path, resumen, huecos)
    print(json.dumps({
        k: v for k, v in resumen.items()
        if k in ("fecha", "fuente", "similares_unicos", "inventario_activo",
                 "cubiertos", "huecos", "cobertura_pct", "pedido_surtir", "pzas_surtir")
    }, indent=2))
    print("surtir", xlsx_surtir, "renglones", len(pedido))
    print("md", md_path)
    print("xlsx", xlsx_pedido)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
