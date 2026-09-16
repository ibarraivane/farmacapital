#!/usr/bin/env python3
"""Pares curados: ficha pobre sin EAN (carga inicial) → ficha completa que ya existe.

No suma stock al desactivar: en varios el mismo lote se contó en los dos SKUs.
"""
from __future__ import annotations

import csv
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
REPORTE = ROOT / "pricing" / "reportes"
HOY = date.today().isoformat().replace("-", "")

# sku_pobre, nombre_pobre, stock_pobre, sku_bueno, nombre_bueno, ean, stock_bueno, nota
DUPLICADOS = [
    ("FC-A0D320D1", "Amoxicilina 500 mg 12 cápsulas", 5,
     "FC-49021570", "Amoxicilina 500 mg Caja con 12 cápsulas AMSA", "7501349021570", 4,
     "Misma caja AMSA. La ficha pobre es solo el genérico."),
    ("FC-022543CD", "Valclan 500/125 mg 10 tabletas", 3,
     "FC-01007199", "Valclan 500 mg / 125 mg Caja con 10 tabletas Wandel", "7503001007199", 3,
     "Misma presentación 500/125. No confundir con Valclan 875."),
    ("FC-7D1D9857", "Acetilsalicilico 100 mg 30 tabletas", 5,
     "FC-42803524", "Ácido acetilsalicílico 100 mg Caja con 30 tabletas beadvance", "7501342803524", 5,
     "Es AAS 100 mg, no Aspirina 500. El EAN de Aspirina ya se le había quitado."),
    ("FC-95779436", "Ácido acetilsalicílico efervescente 300 mg C/20", 5,
     "EQ-ALP0300", "Acido Acetilsalicilico Ef 20 Tab 300 Mg", "7501384504908", 5,
     "Misma caja Psicofarma 300 mg × 20."),
    ("FC-B25B4654", "Cina (Ciprofloxacino) 750 mg 7 tabletas", 2,
     "FC-52200809", "Cina 750 mg Caja con 7 tabletas Landsteiner", "7502225092486", 2,
     "Cina es levofloxacino; la ficha pobre puso ciprofloxacino."),
    ("FC-6C2878CF", "Budenova budesonida 0.125 mg/ml 5 amp × 2 ml", 1,
     "EQ-NOV165", "Budenova Susp .125 Mg/Ml 5 Amp 2ml", "7501075726251", 1,
     "Misma Budenova Novag."),
    ("FC-26EA40A4", "Raamcinet cetirizina 10 mg C/10", 6,
     "FC-27872123", "Raamcinet cetirizina 10 mg", "7502227872123", 4,
     "Misma caja C/10. Stocks 6 y 4: probable doble conteo."),
    ("FC-44B6751A", "LAÜR Adulto solución inyectable C/3", 1,
     "EQ-SON264", "Laur Solución Inyectable", "7502001166981", 1,
     "Adulto. El infantil (7502001167001) es otro SKU."),
    ("FC-1321B34F", "Hidroxin hidroxizina 10 mg C/30", 1,
     "EQ-MAI099", "Hidroxin", "785120754681", 1,
     "Misma Hidroxin Mavi 10 mg C/30."),
    ("FC-AA7B0686", "Drosequim Adulto jarabe 300/160 mg 200 ml", 1,
     "EQ-QUM070", "Drosequim Ad 1 Jbe 300/160/200 Ml", "7502223111400", 1,
     "Mismo jarabe adulto Quimpharma."),
    ("FC-926099D3", "Merthiolate Rojo Kohn 20 ml", 5,
     "FC-46601138", "Merthiolate Rojo Kohn 20 ml", "7506346601138", 5,
     "Mismo nombre. Stocks 5 y 5: probable doble conteo."),
    ("EQ-MAV401", "Dexpantenol 1 Cma 5% 30 G", 2,
     "FC-09749421", "Dexpantenol crema 5% 30 g (Maver Tattoo)", "7502009749421", 2,
     "Misma crema Maver 5% 30 g. Pamedan es otra marca."),
    ("EQ-BRL072-1", "Lo Bruquin 2 Tab 150/200 Mg", 10,
     "EQ-BRL072", "Lo Bruquin Quinfamida, Albendazol", "7502208894915", 10,
     "SKU -1 sin EAN. Ambos con stock 10: el mismo lote contado dos veces."),
    ("FMX-501619", "Eucalin-Miel Jarabe C/120 Ml", 1,
     "FC-08100013", "Eucalín jarabe miel y propóleo 120 ml", "714908100013", 1,
     "Carga Farma MX vs ficha completa."),
    ("FMX-502465", "Colageno-Naturex Tabletas C/60", 1,
     "FC-9741524", "Naturex colágeno hidrolizado tabletas C/60", "7502009741524", 1,
     "Carga Farma MX vs ficha Naturex."),
    ("FMX-301136", "Cintapore cinta microporosa piel 2.5 cm x 5 m", 2,
     "FC-84500546", "Cintapore cinta microporosa piel 2.5 cm x 5m", "7506484500546", 0,
     "El de EAN tiene stock 0: pasar las 2 piezas a ese SKU."),
    ("FMX-501000", "Gelcavit-Colors Capsulas C/30", 1,
     "FC-30713547", "Gelcavit Colors capsulas 0.62 g C/30", "7501130713547", 1,
     "Misma variante Colors."),
    ("FMX-500998", "Gelcavit-Platinum Capsulas C/30", 1,
     "FC-30713851", "Gelcavit Platinum capsulas 1.39 g C/30", "7501130713851", 1,
     "Misma variante Platinum."),
    ("FMX-501003", "Gelcavit-Q-10 Capsulas C/30", 1,
     "FC-13071164", "Gelcavit Q-10 coenzima Q10 jalea real C/30", "7501130711642", 1,
     "Misma variante Q-10."),
    ("FMX-505937", "Pleniform-40 Tabletas C/30", 1,
     "FC-1041884", "Biomiral Pleniform 40 isoflavonas soya C/30", "7503181041884", 1,
     "Mismo Pleniform 40."),
    ("FMX-302947", "Citrato Magnesio/Lecitina Soya-Naturex Capsulas C/30", 1,
     "FC-9892403", "Naturex citrato de magnesio y lecitina de soya C/30", "7502259892403", 2,
     "Misma Naturex citrato + lecitina."),
    ("FMX-506935", "Normogotero-Sensimedical Piezas C/1 S/Aguja", 5,
     "FC-22322395", "Normogotero Sensi Medical", "7506022322395", 5,
     "Mismo gotero. Stocks 5 y 5: probable doble conteo."),
]

# Sin EAN pero NO son el mismo producto que otro SKU (les falta código, no desactivar).
UNICOS_SIN_EAN = [
    ("FC-DB4A39AE", "Eferox (Cefalexina) 500 mg 12 comprimidos", 1, "No hay otra Eferox. Solo falta EAN."),
    ("FC-63975795", "Gentamicina 1 mg 25 comprimidos", 3, "No es la gentamicina inyectable 160 mg AMSA."),
    ("FC-262F2A30", "Irbesartan 150 mg 14 tabletas AMSA", 3, "Los LGEN son 28 tabletas; AMSA 300 mg es otra dosis."),
    ("FC-2E5B7248", "Reumatol 60 g", 2, "Único. Sin pareja con EAN."),
    ("FC-3E863E37", "Tratidri gel 500/50 mg 60 g", 1, "Único."),
    ("FC-405A75E3", "Ursodesoxicolico 250 mg 50 cápsulas", 1, "Único."),
    ("EQ-SON225", "Norkin 7 Caps 40 Mg", 2, "Único."),
    ("EQ-AMS234", "Pregabalina 150 mg C/28 cápsulas AMSA", 1, "Hay 75 mg con EAN; esta dosis 150 no."),
    ("FC-C22EBFE6", "Perilla N2", 2, "Tamaño distinto a la N1 (sí tiene EAN)."),
    ("FC-614E4F82", "Perilla N3", 2, "Tamaño distinto a la N1."),
    ("FC-FFC25DD1", "Perilla N4", 2, "Tamaño distinto a la N1."),
    ("FC-A871D831", "Perilla N6", 2, "Tamaño distinto a la N1."),
    ("FC-9A1C64E7", "Perilla Edigar N O Caja", 2, "Revisar si es caja o pieza; no es la N1."),
    ("FC-C8B741F6", "FC producto botiquín", 2, "Ficha vacía. Identificar qué pieza es."),
    ("FMX-301721", "Vita/Kid/C Jarabe C/240 Ml", 1, "Único."),
    ("FMX-302884", "Sol-Sun Crema C/50 Gr 50-Fps", 2, "Único."),
    ("FMX-502700", "La-Femme Capsulas C/30", 1, "Único."),
    ("FMX-301138", "Cinta-Microporosa-Codifarma 2.5 cm × 5 m blanco", 17, "Otra marca (Codifarma), no Cintapore."),
    ("FC-IFC-83368", "Gel sanitizante Dibar 50 ml", 5, "Único. Falta EAN."),
    ("FC-IFC-83613", "Venditas adhesivas redondas Jayor C/100", 1, "Único. Falta EAN."),
    ("FC-IFC-83490", "Guantes de nitrilo azul chico C/100", 1, "Único. Falta EAN."),
    ("FC-IFC-83125", "Guantes de nitrilo azul grande C/100", 1, "Único. Falta EAN."),
    ("FC-IFC-83947", "Guantes de nitrilo negro mediano C/100", 1, "Único. Falta EAN."),
    ("FC-IFC-82084", "Brocha para tinte con peine de cola", 4, "Único. Falta EAN."),
    ("FC-IFC-83552", "Venda Stick cohesiva 2 pulg × 4.5 m rojo", 2, "Único. Falta EAN."),
    ("FC-IFC-82912A", "Venda Stick cohesiva 3 pulg × 4.5 m azul", 1, "Único. Falta EAN."),
    ("FC-IFC-82912P", "Venda Stick cohesiva 3 pulg × 4.5 m piel", 1, "Único. Falta EAN."),
    ("FC-EXP-PALM8", "Palmolive Neutro Balance jabón 100 g 8 pack", 1, "Único. Falta EAN."),
    ("FC-89F00320", "Mercurio Arnica C/25", 5, "Parece Árnica Mercurio glóbulos, pero no lo desactivo sin ver la caja."),
    ("FC-DFF99C3F", "Mercurio (ficha mezclada)", 3, "La ficha está sucia (árnica + jarabe de granada). Identificar en anaquel."),
]

# Huecos reales de mostrador: no están ni como ficha pobre ni con EAN.
COMPRAR = [
    (10, "Metformina 1000 mg 30 tabletas", "Diabetes — hoy solo hay 500 y 850"),
    (6, "Metformina 750 mg LP 30 tabletas", "Diabetes liberación prolongada"),
    (6, "Racecadotrilo 30 mg 18 sobres", "Diarrea niños — no hay Hidrasec ni genérico"),
    (6, "Racecadotrilo 10 mg 18 sobres", "Diarrea bebés"),
    (3, "Racecadotrilo 100 mg 9 cápsulas", "Diarrea adulto"),
    (6, "Atenolol 50 mg 28 tabletas", "Presión — no hay betabloqueador"),
    (6, "Atenolol 100 mg 28 tabletas", "Presión"),
]


def escribir_xlsx(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    fill_h = PatternFill("solid", fgColor="1A1A1A")
    font_h = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    thin = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )
    fill_alt = PatternFill("solid", fgColor="FFF8E7")
    fill_pass = PatternFill("solid", fgColor="D5F5E3")

    def header(ws, headers):
        for i, h in enumerate(headers, 1):
            c = ws.cell(1, i, h)
            c.fill = fill_h
            c.font = font_h
            c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 24
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    res = wb.active
    res.title = "Resumen"
    res["A1"] = "SKUs sin código de barras que ya existen (carga inicial)"
    res["A1"].font = Font(name="Calibri", bold=True, size=16)
    res.merge_cells("A1:B1")
    res["A2"] = (
        "Son las fichas de agosto (hex FC-… / FMX / EQ-…) con poco o ningún dato. "
        "La pieza ya está en el catálogo con EAN y nombre de mostrador. "
        "Desactiva la ficha pobre. No sumes stock: en varios el mismo lote se contó dos veces "
        "(Lo Bruquin 10+10, Merthiolate 5+5, Raamcinet 6+4, Normogotero 5+5)."
    )
    res.merge_cells("A2:B5")
    res["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    filas = [
        ("Fecha", date.today().isoformat()),
        ("Duplicados confirmados (SQL desactiva el sin EAN)", len(DUPLICADOS)),
        ("Sin EAN pero únicos (no desactivar)", len(UNICOS_SIN_EAN)),
        ("A comprar — no están en catálogo", len(COMPRAR)),
        ("Único caso para pasar stock", "Cintapore FMX-301136 → FC-84500546 (el de EAN tiene 0)"),
    ]
    res["A7"] = "Métrica"
    res["B7"] = "Valor"
    for col in range(1, 3):
        res.cell(7, col).fill = fill_h
        res.cell(7, col).font = font_h
    for i, (a, b) in enumerate(filas, start=8):
        res.cell(i, 1, a)
        res.cell(i, 2, b)
    res.column_dimensions["A"].width = 64
    res.column_dimensions["B"].width = 56

    ws = wb.create_sheet("DESACTIVAR_ya_existen")
    headers = [
        "Acción", "SKU sin EAN (quitar)", "Nombre pobre", "Stock pobre",
        "SKU que se queda", "Nombre completo", "EAN", "Stock del bueno",
        "¿Pasar stock?", "Nota",
    ]
    header(ws, headers)
    for r_i, row in enumerate(DUPLICADOS, start=2):
        pobre, nom_p, stk_p, bueno, nom_b, ean, stk_b, nota = row
        pasar = "SÍ — el de EAN tiene 0" if stk_b == 0 else "NO — no sumar (posible doble conteo)"
        vals = ["Desactivar ficha pobre", pobre, nom_p, stk_p, bueno, nom_b, ean, stk_b, pasar, nota]
        for c_i, v in enumerate(vals, 1):
            cell = ws.cell(r_i, c_i, v)
            cell.border = thin
            if r_i % 2 == 0:
                cell.fill = fill_alt
            if c_i == 9 and stk_b == 0:
                cell.fill = fill_pass
    widths = [22, 16, 44, 12, 16, 52, 16, 14, 36, 56]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A1:J{len(DUPLICADOS)+1}"

    ws3 = wb.create_sheet("NO_ESTAN_COMPRAR", 1)
    header(ws3, ["Pedir", "Qué comprar (genérico)", "Para qué"])
    for r_i, row in enumerate(COMPRAR, start=2):
        for c_i, v in enumerate(row, 1):
            cell = ws3.cell(r_i, c_i, v)
            cell.border = thin
            if r_i % 2 == 0:
                cell.fill = fill_alt
        ws3.cell(r_i, 1).font = Font(name="Calibri", bold=True, size=14)
    for i, w in enumerate([8, 48, 48], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("UNICOS_falta_EAN")
    header(ws2, ["SKU", "Nombre", "Stock", "Por qué no es duplicado"])
    for r_i, row in enumerate(UNICOS_SIN_EAN, start=2):
        for c_i, v in enumerate(row, 1):
            cell = ws2.cell(r_i, c_i, v)
            cell.border = thin
            if r_i % 2 == 0:
                cell.fill = fill_alt
    for i, w in enumerate([16, 52, 10, 56], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def escribir_csv(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow([
            "accion", "sku_pobre", "nombre_pobre", "stock_pobre",
            "sku_bueno", "nombre_bueno", "ean", "stock_bueno", "pasar_stock", "nota",
        ])
        for row in DUPLICADOS:
            pobre, nom_p, stk_p, bueno, nom_b, ean, stk_b, nota = row
            pasar = "si" if stk_b == 0 else "no"
            w.writerow(["desactivar", pobre, nom_p, stk_p, bueno, nom_b, ean, stk_b, pasar, nota])


def _sql_values() -> str:
    rows = []
    for pobre, nom_p, stk_p, bueno, nom_b, ean, stk_b, nota in DUPLICADOS:
        pasar = "true" if stk_b == 0 else "false"
        def q(s: str) -> str:
            return "'" + s.replace("'", "''") + "'"
        rows.append(
            f"  ({q(pobre)}, {q(nom_p)}, {q(bueno)}, {q(ean)}, {pasar})"
        )
    return ",\n".join(rows)


def escribir_sql(path: Path) -> None:
    values = _sql_values()
    sql = f"""-- Desactiva fichas pobres SIN EAN que ya existen con código de barras.
-- 22 pares confirmados. No toca Mercurio Árnica (revisar caja).
--
-- Candados:
--   · el SKU pobre sigue activo y sin codigo_barras
--   · el SKU bueno está activo y tiene exactamente ese EAN
--   · NO suma stock (el mismo lote se contó dos veces), salvo Cintapore
--     (FMX-301136 → FC-84500546: el de EAN tiene 0; pasa las 2 piezas)
--   · deja stock 0 y apaga lotes del SKU pobre para que no siga en caducidad
--
-- Idempotente. Correr en Supabase SQL Editor.

begin;

create temporary table _fc_dup_sin_ean (
  sku_pobre text primary key,
  nombre_pobre text,
  sku_bueno text not null,
  ean_bueno text not null,
  pasar_stock boolean not null default false
) on commit drop;

insert into _fc_dup_sin_ean (sku_pobre, nombre_pobre, sku_bueno, ean_bueno, pasar_stock)
values
{values};

-- Vista previa (debe dar 22 filas listo_para_desactivar).
select
  d.sku_pobre,
  p.nombre as pobre_nombre,
  p.stock as pobre_stock,
  d.sku_bueno,
  b.nombre as bueno_nombre,
  b.codigo_barras as bueno_ean,
  b.stock as bueno_stock,
  d.pasar_stock,
  case
    when p.id is null then 'pobre_ya_no_existe'
    when p.activo is not true then 'pobre_ya_inactivo'
    when nullif(btrim(p.codigo_barras), '') is not null then 'pobre_ahora_tiene_ean'
    when b.id is null then 'bueno_no_existe'
    when b.activo is not true then 'bueno_inactivo'
    when regexp_replace(coalesce(b.codigo_barras, ''), '\\D', '', 'g')
         <> regexp_replace(d.ean_bueno, '\\D', '', 'g') then 'ean_del_bueno_no_cuadra'
    else 'listo_para_desactivar'
  end as estado
from _fc_dup_sin_ean d
left join public.productos p on p.sku = d.sku_pobre
left join public.productos b on b.sku = d.sku_bueno
order by d.sku_pobre;

-- Cintapore: el de EAN tiene 0. Pasa el stock del pobre.
update public.productos b
   set stock = coalesce(b.stock, 0) + coalesce(p.stock, 0)
  from _fc_dup_sin_ean d
  join public.productos p
    on p.sku = d.sku_pobre
   and p.activo = true
   and (p.codigo_barras is null or btrim(p.codigo_barras) = '')
 where d.pasar_stock
   and b.sku = d.sku_bueno
   and b.activo = true
   and regexp_replace(coalesce(b.codigo_barras, ''), '\\D', '', 'g')
       = regexp_replace(d.ean_bueno, '\\D', '', 'g')
   and coalesce(b.stock, 0) = 0;

update public.lotes l
   set activo = false
  from public.productos p
  join _fc_dup_sin_ean d on d.sku_pobre = p.sku
  join public.productos b on b.sku = d.sku_bueno
 where l.producto_id = p.id
   and l.activo = true
   and p.activo = true
   and (p.codigo_barras is null or btrim(p.codigo_barras) = '')
   and b.activo = true
   and regexp_replace(coalesce(b.codigo_barras, ''), '\\D', '', 'g')
       = regexp_replace(d.ean_bueno, '\\D', '', 'g');

update public.productos p
   set activo = false,
       stock = 0
  from _fc_dup_sin_ean d
  join public.productos b on b.sku = d.sku_bueno
 where p.sku = d.sku_pobre
   and p.activo = true
   and (p.codigo_barras is null or btrim(p.codigo_barras) = '')
   and b.activo = true
   and regexp_replace(coalesce(b.codigo_barras, ''), '\\D', '', 'g')
       = regexp_replace(d.ean_bueno, '\\D', '', 'g');

-- Verificación: pobres deben quedar inactivos.
select p.sku, p.nombre, p.activo, p.stock, p.codigo_barras
  from public.productos p
  join _fc_dup_sin_ean d on d.sku_pobre = p.sku
 order by p.sku;

commit;
"""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(sql, encoding="utf-8")


def main() -> int:
    tag = HOY
    xlsx = REPORTE / f"duplicados_sin_barcode_{tag}.xlsx"
    csv_p = REPORTE / f"duplicados_sin_barcode_{tag}.csv"
    sql_p = ROOT / "sql" / f"patch_desactivar_duplicados_sin_barcode_{tag}.sql"
    escribir_xlsx(xlsx)
    escribir_csv(csv_p)
    escribir_sql(sql_p)
    print(xlsx)
    print(csv_p)
    print(sql_p)
    print("duplicados", len(DUPLICADOS), "unicos_sin_ean", len(UNICOS_SIN_EAN))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
